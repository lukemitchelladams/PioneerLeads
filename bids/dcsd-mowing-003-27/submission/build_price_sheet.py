"""Fill DCSD's own "Mowing Sites Price Sheet.xlsx" (RFP 003-27) from Alex's per-site prices plus Brown Bison's markup.

    python3 bids/dcsd-mowing-003-27/submission/build_price_sheet.py

Input: private/alex_prices.json (git-ignored), Alex's numbers exactly as he sent them:
  {"markup": 0.05,
   "sites": {"Chaparral HS": {"turf_weekly": 0, "rock_weekly": 0, "mulch_weekly": 0, "native_monthly": 0,
                              "monthly_total": 0, "seasonal_total": 0}, ...},
   "additional_mow_per_school": {"Chaparral HS": 0, ...},          # Q&A 1 Q11: priced per school, separate
   "additional_services": [["Spring aeration, all yards", "per season", 0], ...],
   "expected_seasonal_grand_total_with_markup": 409489.07}
Every dollar figure is raised by the markup and rounded to the cent. Group and grand totals are re-summed from the
marked-up site totals and checked against the expected grand total; a mismatch over $1 stops the build.
Output: private/Brown Bison Holdings LLC - DCSD 003-27 Price Sheet.xlsx
"""
from __future__ import annotations

import json
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

BID = Path(__file__).resolve().parents[1]
FORM = BID / "docs" / "Mowing Sites Price Sheet.xlsx"
INPUTS = BID / "private" / "alex_prices.json"
OUT = BID / "private" / "Brown Bison Holdings LLC - DCSD 003-27 Price Sheet.xlsx"
COLS = {"turf_weekly": "D", "rock_weekly": "F", "mulch_weekly": "H", "native_monthly": "J",
        "monthly_total": "K", "seasonal_total": "L"}
GROUP_TOTAL_ROWS = {"GROUP 2 TOTALS", "GROUP 5 TOTALS", "GROUP 7 TOTALS", "GROUP 8 TOTALS"}
MONEY = '"$"#,##0.00'


def up(x, markup) -> Decimal:
    return (Decimal(str(x)) * (Decimal(1) + Decimal(str(markup)))).quantize(Decimal("0.01"), ROUND_HALF_UP)


def main() -> int:
    cfg = json.loads(INPUTS.read_text())
    m = cfg["markup"]
    wb = openpyxl.load_workbook(FORM)
    ws = wb.active
    seen, group_k, group_l, grand_k, grand_l = set(), Decimal(0), Decimal(0), Decimal(0), Decimal(0)
    for row in range(1, ws.max_row + 1):
        name = (ws[f"A{row}"].value or "").strip()
        if name in cfg["sites"]:
            site = cfg["sites"][name]
            for key, col in COLS.items():
                if site.get(key) is None:
                    sys.exit(f"{name}: {key} missing")
                v = up(site[key], m)
                ws[f"{col}{row}"] = float(v)
                ws[f"{col}{row}"].number_format = MONEY
                if key == "monthly_total":
                    group_k += v
                if key == "seasonal_total":
                    group_l += v
            seen.add(name)
        elif name in GROUP_TOTAL_ROWS:
            ws[f"K{row}"], ws[f"L{row}"] = float(group_k), float(group_l)
            ws[f"K{row}"].number_format = ws[f"L{row}"].number_format = MONEY
            grand_k, grand_l, group_k, group_l = grand_k + group_k, grand_l + group_l, Decimal(0), Decimal(0)
        elif str(ws[f"J{row}"].value or "").strip() == "Grand Totals":
            ws[f"K{row}"], ws[f"L{row}"] = float(grand_k), float(grand_l)
            ws[f"K{row}"].number_format = ws[f"L{row}"].number_format = MONEY
    missing = set(cfg["sites"]) ^ seen
    if missing:
        sys.exit(f"site names that did not match the District form: {sorted(missing)}")
    expect = Decimal(str(cfg["expected_seasonal_grand_total_with_markup"]))
    if abs(grand_l - expect) > 1:
        sys.exit(f"seasonal grand total {grand_l} does not match the expected {expect}")

    extra = wb.create_sheet("Additional Pricing")
    extra.append(["Additional mow (per school, per occurrence) - Q&A 1, Q11", "Price"])
    for school, price in cfg.get("additional_mow_per_school", {}).items():
        extra.append([school, float(up(price, m))])
    extra.append([])
    extra.append(["Additional services", "Unit", "Price"])
    for desc, unit, price in cfg.get("additional_services", []):
        extra.append([desc, unit, float(up(price, m))])
    for row in extra.iter_rows():
        for c in row:
            if isinstance(c.value, float):
                c.number_format = MONEY
    extra.column_dimensions["A"].width = 60
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(BID.parents[1])}: seasonal grand total ${grand_l:,.2f} (expected ${expect:,.2f})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
